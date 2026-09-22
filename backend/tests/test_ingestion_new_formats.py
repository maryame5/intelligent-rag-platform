import io

from docx import Document as DocxDocument

from app.models.chunk import Chunk


def register_and_login(client, email, password="strongpass123"):
    client.post("/auth/register", json={"email": email, "password": password})
    login = client.post("/auth/login", json={"email": email, "password": password})
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def _build_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Politique de congés", level=1)
    doc.add_paragraph("chat chat chat 25 jours par an.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_upload_is_ingested_with_section_metadata(
    client, fake_minio, worker_session_override, fake_embeddings, db_session
):
    headers = register_and_login(client, "docxuser@example.com")
    kb = client.post("/knowledge-bases", json={"name": "Docs"}, headers=headers).json()

    docx_content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    upload = client.post(
        f"/knowledge-bases/{kb['id']}/documents",
        files={"file": ("politique.docx", _build_docx_bytes(), docx_content_type)},
        headers=headers,
    )
    assert upload.status_code == 202

    job_id = upload.json()["job_id"]
    job = client.get(f"/jobs/{job_id}", headers=headers)
    assert job.json()["status"] == "SUCCESS"

    document_id = upload.json()["document"]["id"]
    chunks = db_session.query(Chunk).filter(Chunk.document_id == document_id).all()
    assert len(chunks) >= 1
    assert chunks[0].section == "Politique de congés"


def test_html_upload_is_ingested_with_section_metadata(
    client, fake_minio, worker_session_override, fake_embeddings, db_session
):
    headers = register_and_login(client, "htmluser@example.com")
    kb = client.post("/knowledge-bases", json={"name": "Docs"}, headers=headers).json()

    html_content = b"<html><body><h1>Introduction</h1><p>chat chat chat contenu html.</p></body></html>"
    upload = client.post(
        f"/knowledge-bases/{kb['id']}/documents",
        files={"file": ("page.html", html_content, "text/html")},
        headers=headers,
    )
    assert upload.status_code == 202

    job_id = upload.json()["job_id"]
    job = client.get(f"/jobs/{job_id}", headers=headers)
    assert job.json()["status"] == "SUCCESS"

    document_id = upload.json()["document"]["id"]
    chunks = db_session.query(Chunk).filter(Chunk.document_id == document_id).all()
    assert len(chunks) >= 1
    assert chunks[0].section == "Introduction"


def test_markdown_upload_is_ingested_with_section_metadata(
    client, fake_minio, worker_session_override, fake_embeddings, db_session
):
    headers = register_and_login(client, "mduser@example.com")
    kb = client.post("/knowledge-bases", json={"name": "Docs"}, headers=headers).json()

    md_content = b"# Notes\nchat chat chat contenu markdown."
    upload = client.post(
        f"/knowledge-bases/{kb['id']}/documents",
        files={"file": ("notes.md", md_content, "text/markdown")},
        headers=headers,
    )
    assert upload.status_code == 202

    job_id = upload.json()["job_id"]
    job = client.get(f"/jobs/{job_id}", headers=headers)
    assert job.json()["status"] == "SUCCESS"

    document_id = upload.json()["document"]["id"]
    chunks = db_session.query(Chunk).filter(Chunk.document_id == document_id).all()
    assert len(chunks) >= 1
    assert chunks[0].section == "Notes"
