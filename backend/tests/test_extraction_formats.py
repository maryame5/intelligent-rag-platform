import io

from docx import Document as DocxDocument

from app.services.extraction import (
    ExtractionError,
    extract_pages,
    extract_pages_from_docx,
    extract_pages_from_html,
    extract_pages_from_markdown,
)


def _build_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Ceci est le paragraphe d'introduction.")
    doc.add_heading("Détails", level=1)
    doc.add_paragraph("Ceci est le paragraphe de détails.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_extraction_splits_by_heading_style():
    data = _build_docx_bytes()
    blocks = extract_pages_from_docx(data)

    sections = [section for _page, section, _text in blocks]
    assert sections == ["Introduction", "Détails"]
    assert "paragraphe d'introduction" in blocks[0][2]
    assert "paragraphe de détails" in blocks[1][2]


def test_docx_extraction_page_is_always_one():
    data = _build_docx_bytes()
    blocks = extract_pages_from_docx(data)
    assert all(page == 1 for page, _section, _text in blocks)


def test_docx_extraction_rejects_corrupted_file():
    try:
        extract_pages_from_docx(b"ceci n'est pas un vrai .docx")
        assert False, "aurait dû lever ExtractionError"
    except ExtractionError:
        pass


def test_markdown_extraction_splits_by_headings():
    text = "# Introduction\nTexte intro.\n\n## Sous-partie\nTexte sous-partie.\n"
    blocks = extract_pages_from_markdown(text.encode("utf-8"))

    sections = [section for _page, section, _text in blocks]
    assert sections == ["Introduction", "Sous-partie"]


def test_markdown_extraction_keeps_preamble_before_first_heading():
    text = "Texte avant tout titre.\n\n# Titre\nTexte du titre.\n"
    blocks = extract_pages_from_markdown(text.encode("utf-8"))

    assert blocks[0][1] is None  # préambule sans section
    assert "avant tout titre" in blocks[0][2]
    assert blocks[1][1] == "Titre"


def test_markdown_without_headings_returns_single_block():
    text = "Juste du texte, aucun titre."
    blocks = extract_pages_from_markdown(text.encode("utf-8"))
    assert len(blocks) == 1
    assert blocks[0][1] is None


def test_html_extraction_splits_by_heading_tags():
    html = """
    <html><body>
        <h1>Introduction</h1>
        <p>Texte d'intro.</p>
        <h2>Details</h2>
        <p>Texte de details.</p>
    </body></html>
    """.encode("utf-8")
    blocks = extract_pages_from_html(html)

    sections = [section for _page, section, _text in blocks]
    assert sections == ["Introduction", "Details"]
    assert "Texte d'intro" in blocks[0][2]


def test_html_extraction_strips_script_and_style_tags():
    html = b"""
    <html><body>
        <style>.hidden { display: none; }</style>
        <script>alert('should not appear');</script>
        <p>Contenu visible.</p>
    </body></html>
    """
    blocks = extract_pages_from_html(html)
    full_text = " ".join(text for _page, _section, text in blocks)
    assert "should not appear" not in full_text
    assert "display: none" not in full_text
    assert "Contenu visible" in full_text


def test_html_without_structure_falls_back_to_raw_text():
    html = b"<html><body>Juste du texte flottant, sans balises p ou h1.</body></html>"
    blocks = extract_pages_from_html(html)
    assert len(blocks) == 1
    assert "Juste du texte flottant" in blocks[0][2]


def test_extract_pages_dispatches_by_mime_type():
    assert extract_pages(b"Bonjour", "text/plain")[0][2] == "Bonjour"
    assert extract_pages("# Titre\ntexte".encode(), "text/markdown")[0][1] == "Titre"
    assert extract_pages(b"<p>contenu</p>", "text/html")[0][2] == "contenu"


def test_extract_pages_rejects_unsupported_mime_type():
    try:
        extract_pages(b"peu importe", "application/zip")
        assert False, "aurait dû lever ExtractionError"
    except ExtractionError:
        pass
